import datetime
import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.borders import Border, Side

from app.db.models import Quote

# Đúng cấu trúc sheet "2- LINH KIỆN BÁN SẴN -ENG" trong file mẫu
# "01. Tài liệu tham khảo/Form báo giá.xlsx" — cột B..P, tiêu đề 3 dòng (5-7), dữ liệu từ dòng 8.
_THIN = Side(style="thin")
_BORDER_ALL = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
_CURRENCY_FMT = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
_YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")

# Chủng loại -> cột đánh dấu "●", khớp app.crud.parts.LOAI_OPTIONS
_LOAI_COLUMN = {
    "Điện": "C",
    "Khí nén": "D",
    "Thủy lực": "E",
    "Cơ khí": "F",
    "Khác": "G",
}

_COLUMN_WIDTHS = {
    "A": 1, "B": 5, "C": 5, "D": 5, "E": 5, "F": 5, "G": 5,
    "H": 53, "I": 35, "J": 14.5, "K": 10, "L": 22.7, "M": 7, "N": 5, "O": 13, "P": 13,
}


def _center(wrap: bool = False) -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def build_parts_excel(quote: Quote) -> io.BytesIO:
    """Xuất danh sách linh kiện tiêu chuẩn của 1 báo giá ra Excel, theo đúng mẫu công ty."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sach linh kien"

    for col, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C8"

    ws["P1"] = "Mã số: BM-0401.09-QLSX / Ngày hiệu lực: 25-Feb-2018 / Lần sửa đổi:00"

    ws["B3"] = "BẢNG TỔNG HỢP LINH KIỆN BÁN SẴN"
    ws["B3"].font = Font(name="Arial", size=16, bold=True)
    ws.row_dimensions[3].height = 21

    du_an = quote.project.ten_thiet_bi if quote.project else (quote.ten_template or "")
    khach_hang = (quote.project.khach_hang if quote.project else quote.khach_hang) or "—"
    ws["B4"] = (
        f"Dự án: {quote.project.ma_du_an if quote.project else '—'} - {du_an}"
        f"  |  Khách hàng: {khach_hang}  |  Mã báo giá: {quote.ma_bao_gia}"
        f"  |  Ngày xuất: {datetime.date.today().strftime('%d-%m-%Y')}"
    )
    ws["B4"].font = Font(name="Arial", size=10, italic=True)

    headers = {
        "B5": "No", "C5": "Chủng loại", "H5": "Tên linh kiện", "I5": "Mã linh kiện",
        "J5": "Maker", "K5": "Xuất xứ", "L5": "Nhà cung cấp", "M5": "Đơn vị",
        "N5": "SL", "O5": "Đơn giá\n(VND)", "P5": "Thành tiền\n(VND)",
    }
    for coord, text in headers.items():
        cell = ws[coord]
        cell.value = text
        cell.font = Font(name="Arial", size=11)
        cell.alignment = _center(wrap=True)
        cell.border = _BORDER_ALL
    for rng in ("B5:B7", "C5:G5", "H5:H7", "I5:I7", "J5:J7", "K5:K7", "L5:L7", "M5:M7", "N5:N7", "O5:O7", "P5:P7"):
        ws.merge_cells(rng)

    sub_labels = {"C6": "Điện", "D6": "Khí nén", "E6": "Thủy\n lực", "F6": "Cơ khí", "G6": "Khác"}
    for coord, text in sub_labels.items():
        cell = ws[coord]
        cell.value = text
        cell.font = Font(name="Arial", size=9)
        cell.alignment = _center(wrap=True)
        cell.border = _BORDER_ALL
    for col in "CDEFG":
        cell = ws[f"{col}7"]
        cell.value = "●"
        cell.font = Font(name="Arial", size=9)
        cell.alignment = _center(wrap=True)
        cell.border = _BORDER_ALL
    ws.row_dimensions[5].height = 14.7
    ws.row_dimensions[6].height = 38.25
    ws.row_dimensions[7].height = 7.2

    row = 8
    no = 1

    def write_header_row(ten: str) -> None:
        nonlocal row, no
        ws[f"B{row}"] = no
        ws[f"H{row}"] = ten
        ws[f"P{row}"] = f"=O{row}*N{row}"
        for col in "BCDEFGHIJKLMNOP":
            cell = ws[f"{col}{row}"]
            cell.border = _BORDER_ALL
            cell.fill = _YELLOW_FILL
            cell.font = Font(name="Arial", size=11)
        ws[f"B{row}"].alignment = _center()
        ws[f"H{row}"].alignment = Alignment(horizontal="left")
        ws[f"O{row}"].number_format = _CURRENCY_FMT
        ws[f"P{row}"].number_format = _CURRENCY_FMT
        row += 1
        no += 1

    def write_item_row(line: "object") -> None:
        nonlocal row, no
        part = line.part
        loai = part.loai if part else None
        xuat_xu = part.xuat_xu if part else None

        ws[f"B{row}"] = no
        if loai in _LOAI_COLUMN:
            ws[f"{_LOAI_COLUMN[loai]}{row}"] = "●"
        ws[f"H{row}"] = line.ten_linh_kien
        ws[f"I{row}"] = line.ma_linh_kien or ""
        ws[f"J{row}"] = line.hang or ""
        ws[f"K{row}"] = xuat_xu or ""
        ws[f"L{row}"] = line.nha_cung_cap or ""
        ws[f"M{row}"] = line.don_vi or ""
        ws[f"N{row}"] = line.so_luong
        ws[f"O{row}"] = line.don_gia
        ws[f"P{row}"] = f"=O{row}*N{row}"

        for col in "BCDEFGHIJKLMNOP":
            cell = ws[f"{col}{row}"]
            cell.border = _BORDER_ALL
            cell.font = Font(name="Arial", size=11)
            if col in ("O", "P"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = _center()
        ws[f"O{row}"].number_format = _CURRENCY_FMT
        ws[f"P{row}"].number_format = _CURRENCY_FMT
        row += 1
        no += 1

    for group in quote.groups:
        write_header_row(group.ten_cum)
        for line in group.part_lines:
            write_item_row(line)

    ungrouped = [l for l in quote.part_lines if l.group_id is None]
    if ungrouped:
        write_header_row("(Chưa phân cụm)")
        for line in ungrouped:
            write_item_row(line)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
