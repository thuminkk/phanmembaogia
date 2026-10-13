import io
import re

import openpyxl

FIELD_HEADER_PREFIXES = {
    "ten_linh_kien": ["tên linh kiện"],
    "ma_linh_kien": ["mã linh kiện"],
    "hang": ["maker"],
    "xuat_xu": ["xuất xứ"],
    "nha_cung_cap": ["nhà cung cấp"],
    "don_vi": ["đơn vị"],
    "don_gia": ["đơn giá"],
}

MAX_HEADER_SCAN_ROWS = 20


def _normalize(text) -> str:
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip().casefold()


def list_sheets(file_bytes: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    return wb.sheetnames


def _find_header(ws) -> tuple[int, dict[str, int]] | None:
    for row_idx in range(1, MAX_HEADER_SCAN_ROWS + 1):
        row_values = [ws.cell(row_idx, c).value for c in range(1, ws.max_column + 1)]
        normalized = [_normalize(v) for v in row_values]
        if not any(v.startswith("tên linh kiện") for v in normalized):
            continue

        columns: dict[str, int] = {}
        for field, prefixes in FIELD_HEADER_PREFIXES.items():
            for col_idx, cell_value in enumerate(normalized, start=1):
                if any(cell_value.startswith(p) for p in prefixes):
                    columns[field] = col_idx
                    break
        if "ten_linh_kien" in columns:
            return row_idx, columns
    return None


def parse_sheet(file_bytes: bytes, sheet_name: str) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb[sheet_name]

    header = _find_header(ws)
    if header is None:
        raise ValueError(
            "Không tìm thấy dòng tiêu đề có cột 'Tên linh kiện' trong sheet này — "
            "kiểm tra lại đúng sheet dạng 'Bảng tổng hợp linh kiện'."
        )
    header_row, columns = header

    rows = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        ten = ws.cell(row_idx, columns["ten_linh_kien"]).value if "ten_linh_kien" in columns else None
        if not ten or not str(ten).strip():
            continue

        def get(field):
            col = columns.get(field)
            if not col:
                return None
            val = ws.cell(row_idx, col).value
            if val is None:
                return None
            return str(val).strip() or None

        don_gia = get("don_gia")
        try:
            don_gia = float(don_gia) if don_gia not in (None, "") else None
        except (TypeError, ValueError):
            don_gia = None

        rows.append(
            {
                "ten_linh_kien": str(ten).strip(),
                "ma_linh_kien": get("ma_linh_kien"),
                "hang": get("hang"),
                "xuat_xu": get("xuat_xu"),
                "nha_cung_cap": get("nha_cung_cap"),
                "don_vi": get("don_vi"),
                "don_gia": don_gia,
            }
        )
    return rows
